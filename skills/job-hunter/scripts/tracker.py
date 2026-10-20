"""Excel job tracker: merge search results, keep the user's own edits, track applications.

Sheets: "İlanlar" (ranked, verified jobs), "Elenenler" (rejected: fake / ghost / poor fit, with reason),
"Takip" (applications and follow-up dates), "Aramalar" (search log).

Usage:
  python tracker.py update  tracker.xlsx results.json --label "Yaz stajı / ML / TR+EU"
  python tracker.py seen    tracker.xlsx                 # ids already in the file (agents skip these)
  python tracker.py apply   tracker.xlsx <job id> [--days 7]   # mark applied, set follow-up date
  python tracker.py status  tracker.xlsx <job id> "Mülakat"    # Mülakat / Ret / Teklif / Vazgeçildi
  python tracker.py due     tracker.xlsx                 # follow-ups due today or earlier (JSON)

results.json: list of normalized job records (see common.job) plus the agents' fields:
  score, fit, potential, verification ("Doğrulandı" | "Muhtemel" | "Şüpheli" | "Sahte"),
  blocker (short text, e.g. "Vize: maaş eşiği"; empty if none), verification_note, pros, cons,
  salary_text, visa, rejected_reason (for Elenenler: "Sahte", "Kapandı", ...).
"""

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import utf8_stdout

JOB_COLS = [  # (header, record key, width)
    ("Sıra", None, 6), ("Puan", "score", 7), ("Uyum", "fit", 7), ("Potansiyel", "potential", 10),
    ("Doğrulama", "verification", 12), ("Engel", "blocker", 22), ("Pozisyon", "title", 34), ("Şirket", "company", 20),
    ("Konum", "location", 18), ("Uzaktan", "remote", 8), ("Tür", "employment_type", 12),
    ("Maaş", "salary_text", 16), ("Vize", "visa", 14), ("Artılar", "pros", 40), ("Eksiler", "cons", 40),
    ("Doğrulama notu", "verification_note", 36), ("Kaynak", "source", 12), ("Resmî başvuru linki", "apply_url", 40),
    ("İlan tarihi", "posted", 12), ("Bulunma", "found", 11),
    ("Durum", "status", 13), ("Başvuru tarihi", "applied", 13), ("Takip tarihi", "follow_up", 13), ("Notlar", "notes", 30),
    ("ID", "id", 13),
]
USER_KEYS = {"status", "applied", "follow_up", "notes"}  # never overwritten by a new search
REJ_COLS = [("Pozisyon", "title", 34), ("Şirket", "company", 20), ("Neden", "rejected_reason", 18),
            ("Kanıt", "verification_note", 60), ("Kaynak", "source", 12), ("Link", "url", 40), ("Bulunma", "found", 11), ("ID", "id", 13)]
VERIFY_FILL = {"Doğrulandı": "C6EFCE", "Muhtemel": "FFF2CC", "Şüpheli": "F8CBAD", "Sahte": "FF9999"}
VERIFY_ORDER = {"Doğrulandı": 0, "Muhtemel": 1, "Şüpheli": 2}
HEAD_FILL = PatternFill("solid", fgColor="1F3864")


def today():
    return dt.date.today().isoformat()


def as_text(v):
    if isinstance(v, list):
        return "\n".join(f"• {x}" for x in v)
    if isinstance(v, bool):
        return "Evet" if v else "Hayır"
    return "" if v is None else v


def read_sheet(ws, cols):
    keys = [k for _, k, _ in cols]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(c not in (None, "") for c in r):
            rows.append({k: v for k, v in zip(keys, r) if k})
    return rows


def write_sheet(ws, cols, rows, ranked=False):
    ws.delete_rows(1, ws.max_row)
    for i, (h, _, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    for n, rec in enumerate(rows, 2):
        for i, (h, k, _) in enumerate(cols, 1):
            v = n - 1 if (ranked and h == "Sıra") else as_text(rec.get(k))
            c = ws.cell(row=n, column=i, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if k in ("apply_url", "url") and isinstance(v, str) and v.startswith("http"):
                c.hyperlink, c.font = v, Font(color="0563C1", underline="single")
            if k == "verification" and v in VERIFY_FILL:
                c.fill = PatternFill("solid", fgColor=VERIFY_FILL[v])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def load(path):
    p = Path(path)
    if not p.exists():
        wb = Workbook()
        wb.active.title = "İlanlar"
        for name in ("Elenenler", "Takip", "Aramalar"):
            wb.create_sheet(name)
        wb["Aramalar"].append(["Tarih", "Arama", "Bulunan", "Eklenen", "Elenen"])
        return wb
    return load_workbook(p)


def save(wb, path):
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"'{path}' is open in Excel. Close it and run again; nothing was changed.")


def rank(rows):
    """Reachable jobs first (by label, then score); jobs with a hard blocker (visa, required language,
    experience far beyond the candidate) go last whatever their label. In testing, label-first alone put
    verified-but-unreachable jobs above realistic ones."""
    return sorted(rows, key=lambda r: (1 if r.get("blocker") else 0,
                                       VERIFY_ORDER.get(r.get("verification"), 3), -(float(r.get("score") or 0))))


def rebuild_takip(wb, rows):
    ws = wb["Takip"]
    ws.delete_rows(1, ws.max_row)
    ws.append(["Pozisyon", "Şirket", "Durum", "Başvuru tarihi", "Takip tarihi", "Takip zamanı geldi mi?", "Link", "ID"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
    for r in rows:
        if r.get("status") and r.get("status") != "Başvurulmadı":
            due = r.get("follow_up") and str(r["follow_up"]) <= today() and r.get("status") == "Başvuruldu"
            ws.append([r.get("title"), r.get("company"), r.get("status"), r.get("applied"), r.get("follow_up"),
                       "EVET" if due else "", r.get("apply_url"), r.get("id")])


def cmd_update(path, results, label):
    wb = load(path)
    new = json.loads(Path(results).read_text(encoding="utf-8"))
    jobs = {r["id"]: r for r in read_sheet(wb["İlanlar"], JOB_COLS)}
    rejected = {r["id"]: r for r in read_sheet(wb["Elenenler"], REJ_COLS)}
    added = dropped = 0
    for rec in new:
        rec.setdefault("found", today())
        if rec.get("verification") == "Sahte" or rec.get("rejected_reason"):
            rec.setdefault("rejected_reason", "Sahte" if rec.get("verification") == "Sahte" else "Elendi")
            if rec["id"] not in rejected:
                dropped += 1
            rejected[rec["id"]] = {**rejected.get(rec["id"], {}), **rec}
            jobs.pop(rec["id"], None)
            continue
        old = jobs.get(rec["id"])
        if old:
            rec = {**old, **{k: v for k, v in rec.items() if k not in USER_KEYS and v not in (None, "", [])}, "found": old.get("found")}
        else:
            added += 1
            rec.setdefault("status", "Başvurulmadı")
        jobs[rec["id"]] = rec
    rows = rank(list(jobs.values()))
    write_sheet(wb["İlanlar"], JOB_COLS, rows, ranked=True)
    write_sheet(wb["Elenenler"], REJ_COLS, list(rejected.values()))
    rebuild_takip(wb, rows)
    wb["Aramalar"].append([dt.datetime.now().strftime("%Y-%m-%d %H:%M"), label, len(new), added, dropped])
    save(wb, path)
    print(json.dumps({"file": str(Path(path).resolve()), "jobs_total": len(rows), "added": added,
                      "rejected_new": dropped, "rejected_total": len(rejected)}, ensure_ascii=False))


def set_fields(path, job_id, **fields):
    wb = load(path)
    rows = read_sheet(wb["İlanlar"], JOB_COLS)
    hit = next((r for r in rows if r.get("id") == job_id), None)
    if not hit:
        sys.exit(f"No job with id {job_id}")
    hit.update(fields)
    rows = rank(rows)
    write_sheet(wb["İlanlar"], JOB_COLS, rows, ranked=True)
    rebuild_takip(wb, rows)
    save(wb, path)
    print(json.dumps({"id": job_id, **fields}, ensure_ascii=False))


def main():
    utf8_stdout()
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    u = sub.add_parser("update"); u.add_argument("xlsx"); u.add_argument("results"); u.add_argument("--label", default="")
    s = sub.add_parser("seen"); s.add_argument("xlsx")
    a = sub.add_parser("apply"); a.add_argument("xlsx"); a.add_argument("id"); a.add_argument("--days", type=int, default=7)
    st = sub.add_parser("status"); st.add_argument("xlsx"); st.add_argument("id"); st.add_argument("status")
    d = sub.add_parser("due"); d.add_argument("xlsx")
    x = ap.parse_args()
    if x.cmd == "update":
        cmd_update(x.xlsx, x.results, x.label)
    elif x.cmd == "seen":
        wb = load(x.xlsx)
        ids = [r["id"] for r in read_sheet(wb["İlanlar"], JOB_COLS)] + [r["id"] for r in read_sheet(wb["Elenenler"], REJ_COLS)]
        print(json.dumps(ids))
    elif x.cmd == "apply":
        set_fields(x.xlsx, x.id, status="Başvuruldu", applied=today(),
                   follow_up=(dt.date.today() + dt.timedelta(days=x.days)).isoformat())
    elif x.cmd == "status":
        set_fields(x.xlsx, x.id, status=x.status)
    else:
        rows = read_sheet(load(x.xlsx)["İlanlar"], JOB_COLS)
        due = [{k: r.get(k) for k in ("id", "title", "company", "applied", "follow_up", "apply_url")} for r in rows
               if r.get("status") == "Başvuruldu" and r.get("follow_up") and str(r["follow_up"]) <= today()]
        print(json.dumps(due, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
